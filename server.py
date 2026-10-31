import csv, hashlib, io, json, os, sqlite3, tempfile, secrets, time, re
from email import policy
from email.parser import BytesParser
from datetime import datetime, timedelta
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import urlparse

ROOT=os.path.dirname(os.path.abspath(__file__)); DB=os.path.join(ROOT,'finance.db')
COLS=['stat_date','department','manager','promotion_leader','promotion_group','product_name','actual_sales','operating_cost','goods_cost','express_fee','onsite_promotion_fee','pugongying_fee','qianfan_fee','content_comment_fee','dou_plus_fee','props_competitor_fee','qa_fee','member_fee','other_fee','profit']
NUMERIC=COLS[6:]
def db():
 c=sqlite3.connect(DB); c.row_factory=sqlite3.Row; c.execute('''CREATE TABLE IF NOT EXISTS product_profit_daily (id INTEGER PRIMARY KEY, stat_date TEXT NOT NULL, department TEXT, manager TEXT, promotion_leader TEXT, promotion_group TEXT, product_name TEXT NOT NULL, actual_sales REAL DEFAULT 0, operating_cost REAL DEFAULT 0, goods_cost REAL DEFAULT 0, express_fee REAL DEFAULT 0, onsite_promotion_fee REAL DEFAULT 0, pugongying_fee REAL DEFAULT 0, qianfan_fee REAL DEFAULT 0, content_comment_fee REAL DEFAULT 0, dou_plus_fee REAL DEFAULT 0, props_competitor_fee REAL DEFAULT 0, qa_fee REAL DEFAULT 0, member_fee REAL DEFAULT 0, other_fee REAL DEFAULT 0, profit REAL DEFAULT 0, source_sheet TEXT, source_hash TEXT, created_at TEXT, updated_at TEXT)'''); schema=c.execute("SELECT sql FROM sqlite_master WHERE type='table' AND name='product_profit_daily'").fetchone()[0] or ''
 if 'UNIQUE(stat_date' in schema:
  c.execute('ALTER TABLE product_profit_daily RENAME TO product_profit_daily_old')
  c.execute('''CREATE TABLE product_profit_daily (id INTEGER PRIMARY KEY, stat_date TEXT NOT NULL, department TEXT, manager TEXT, promotion_leader TEXT, promotion_group TEXT, product_name TEXT NOT NULL, actual_sales REAL DEFAULT 0, operating_cost REAL DEFAULT 0, goods_cost REAL DEFAULT 0, express_fee REAL DEFAULT 0, onsite_promotion_fee REAL DEFAULT 0, pugongying_fee REAL DEFAULT 0, qianfan_fee REAL DEFAULT 0, content_comment_fee REAL DEFAULT 0, dou_plus_fee REAL DEFAULT 0, props_competitor_fee REAL DEFAULT 0, qa_fee REAL DEFAULT 0, member_fee REAL DEFAULT 0, other_fee REAL DEFAULT 0, profit REAL DEFAULT 0, source_sheet TEXT, source_hash TEXT, created_at TEXT, updated_at TEXT)''')
  c.execute('INSERT INTO product_profit_daily SELECT * FROM product_profit_daily_old'); c.execute('DROP TABLE product_profit_daily_old')
 c.execute('''CREATE TABLE IF NOT EXISTS users (id INTEGER PRIMARY KEY, username TEXT UNIQUE NOT NULL, password_hash TEXT NOT NULL, password_salt TEXT NOT NULL, created_at TEXT NOT NULL, role TEXT NOT NULL DEFAULT 'user')''');
 if 'role' not in [row[1] for row in c.execute('PRAGMA table_info(users)').fetchall()]: c.execute("ALTER TABLE users ADD COLUMN role TEXT NOT NULL DEFAULT 'user'")
 c.execute("UPDATE users SET role='admin' WHERE username='admin'"); c.execute('''CREATE TABLE IF NOT EXISTS sessions (token_hash TEXT PRIMARY KEY, user_id INTEGER NOT NULL, expires_at INTEGER NOT NULL, created_at TEXT NOT NULL)'''); c.commit(); return c

LOGIN_ATTEMPTS={}
SESSION_DAYS=7
def password_hash(password,salt=None):
 salt=salt or secrets.token_hex(16)
 return salt,hashlib.pbkdf2_hmac('sha256',password.encode(),salt.encode(),210000).hex()
def valid_password(password):return isinstance(password,str) and 8<=len(password)<=128
def json_body(handler):
 length=int(handler.headers.get('Content-Length','0') or 0)
 return json.loads(handler.rfile.read(length) or b'{}')
def client_key(handler):return handler.client_address[0]
def session_user(c,handler):
 cookie=handler.headers.get('Cookie','')
 token=next((x.split('=',1)[1] for x in cookie.split(';') if x.strip().startswith('finance_session=')), '')
 if not token:return None
 row=c.execute('SELECT u.* FROM sessions s JOIN users u ON u.id=s.user_id WHERE s.token_hash=? AND s.expires_at>?',(hashlib.sha256(token.encode()).hexdigest(),int(time.time()))).fetchone()
 return row
def is_admin(user): return bool(user and user['role']=='admin')
def admin_metrics(c):
 rows=c.execute('SELECT COUNT(*) AS count, MIN(stat_date) AS start_date, MAX(stat_date) AS end_date FROM product_profit_daily').fetchone()
 daily=[dict(x) for x in c.execute('SELECT stat_date, COUNT(*) AS row_count FROM product_profit_daily GROUP BY stat_date ORDER BY stat_date DESC').fetchall()]
 return {'database_bytes':os.path.getsize(DB) if os.path.exists(DB) else 0,'upload_storage_bytes':0,'row_count':rows['count'] or 0,'start_date':rows['start_date'],'end_date':rows['end_date'],'daily':daily}
def num(x):
 try:
  text=str(x or 0).strip().replace(',','').replace(' ','').replace('，','').replace('￥','').replace('¥','').replace('元','').replace('日元','')
  negative=text.startswith('(') and text.endswith(')')
  text=text.strip('()')
  match=re.search(r'-?\d+(?:\.\d+)?',text)
  if not match:return 0
  value=float(match.group())
  return round(-abs(value) if negative else value,2)
 except:return 0
def excel_date(v):
 if isinstance(v,(int,float)): return (datetime(1899,12,30)+timedelta(days=int(v))).strftime('%Y-%m-%d')
 return str(v or '')
def sample_seed(c):
 return
 if c.execute('SELECT 1 FROM product_profit_daily LIMIT 1').fetchone(): return
 rows=[]
 for date, vals in [('2026-10-26',[88524.88,7928,-16514.88]),('2026-10-27',[94842.94,8310.98,29448.86])]:
  rows.append((date,'演示部门','演示主管','演示组长','',f'演示产品 {date}',vals[0],0,vals[1],0,0,0,0,0,0,0,0,0,0,vals[2],'demo','demo',datetime.now().isoformat(),datetime.now().isoformat()))
 c.executemany('INSERT INTO product_profit_daily VALUES (NULL,'+','.join('?'*24)+')',rows); c.commit()
def payload(c,start_date='',end_date=''):
 clauses=[]; params=[]
 if start_date: clauses.append('stat_date>=?'); params.append(start_date)
 if end_date: clauses.append('stat_date<=?'); params.append(end_date)
 where=(' WHERE '+' AND '.join(clauses)) if clauses else ''
 rows=c.execute('SELECT * FROM product_profit_daily'+where+' ORDER BY stat_date DESC, id',params).fetchall(); d=[dict(r) for r in rows]
 def total(k):return round(sum((r[k] or 0) for r in d),2)
 revenue=total('actual_sales'); cost=total('goods_cost'); expense=round(total('operating_cost')+sum(total(k) for k in ['express_fee','onsite_promotion_fee','pugongying_fee','qianfan_fee','content_comment_fee','dou_plus_fee','props_competitor_fee','qa_fee','member_fee','other_fee']),2); profit=total('profit')
 def group(key):
  out={}
  for r in d:
   g=out.setdefault(r[key] or '未分类',{'name':r[key] or '未分类','revenue':0,'profit':0});g['revenue']+=r['actual_sales'];g['profit']+=r['profit']
  return sorted(out.values(),key=lambda x:x['profit'],reverse=True)
 by_date=[{'stat_date':x['name'],'revenue':round(x['revenue'],2),'profit':round(x['profit'],2)} for x in group('stat_date')]
 by_product=[{'product_name':x['name'],'revenue':round(x['revenue'],2),'profit':round(x['profit'],2)} for x in group('product_name')]
 by_department=[{'department':x['name'],'revenue':round(x['revenue'],2),'profit':round(x['profit'],2)} for x in group('department')]
 return {'meta':{'rows':len(d),'days':len(set(r['stat_date'] for r in d)),'latest_date':max((r['stat_date'] for r in d),default='')},'summary':{'revenue':revenue,'gross_profit':round(revenue-cost,2),'operating_expense':expense,'net_profit':profit,'gross_margin':round((revenue-cost)/revenue*100,1) if revenue else 0,'net_margin':round(profit/revenue*100,1) if revenue else 0},'by_date':by_date,'by_product':by_product,'by_department':by_department,'rows':d}
class Handler(SimpleHTTPRequestHandler):
 def __init__(self,*a,**kw):super().__init__(*a,directory=ROOT,**kw)
 def send_json(self,data,status=200):
  raw=json.dumps(data,ensure_ascii=False).encode();self.send_response(status);self.send_header('Content-Type','application/json; charset=utf-8');origin=self.headers.get('Origin','');self.send_header('Access-Control-Allow-Origin',origin if origin in ('http://127.0.0.1:5173','http://localhost:5173') else 'http://127.0.0.1:5173');self.send_header('Access-Control-Allow-Credentials','true');self.send_header('Vary','Origin');self.send_header('X-Content-Type-Options','nosniff');self.send_header('Content-Length',str(len(raw)));self.end_headers();self.wfile.write(raw)
 def do_OPTIONS(self):self.send_response(204);self.send_header('Access-Control-Allow-Origin',self.headers.get('Origin','http://127.0.0.1:5173'));self.send_header('Access-Control-Allow-Credentials','true');self.send_header('Access-Control-Allow-Methods','GET,POST,OPTIONS');self.send_header('Access-Control-Allow-Headers','Content-Type');self.end_headers()
 def do_GET(self):
  path=urlparse(self.path).path;c=db();sample_seed(c)
  if path=='/api/health':return self.send_json({'ok':True,'version':'1.0.0','database':'sqlite'})
  if path=='/api/auth/status':return self.send_json({'setup_required':c.execute('SELECT COUNT(*) FROM users').fetchone()[0]==0})
  if path=='/api/auth/me':
   user=session_user(c,self)
   return self.send_json({'authenticated':bool(user),'user':{'id':user['id'],'username':user['username'],'role':user['role']} if user else None},200 if user else 401)
  if path.startswith('/api/') and not session_user(c,self):return self.send_json({'error':'请先登录'},401)
  if path=='/api/admin/metrics':
   user=session_user(c,self)
   if not is_admin(user): return self.send_json({'error':'仅管理员可访问'},403)
   from urllib.parse import parse_qs
   query=parse_qs(urlparse(self.path).query); start=query.get('start_date',[''])[0]; end=query.get('end_date',[''])[0]
   result=admin_metrics(c); result['selected_row_count']=c.execute('SELECT COUNT(*) FROM product_profit_daily WHERE stat_date>=? AND stat_date<=?',(start,end)).fetchone()[0] if start and end else result['row_count']
   return self.send_json(result)
  if path=='/api/summary':
   from urllib.parse import parse_qs
   query=parse_qs(urlparse(self.path).query); start=query.get('start_date',[''])[0]; end=query.get('end_date',['']); end=end[0] if end else ''
   return self.send_json(payload(c,start,end))
  if path=='/api/profit-rows':
   fields='stat_date,department,manager,promotion_leader,product_name,actual_sales,operating_cost,goods_cost,other_fee,profit'
   return self.send_json({'rows':[dict(r) for r in c.execute(f'SELECT {fields} FROM product_profit_daily ORDER BY stat_date DESC,id').fetchall()]})
  return super().do_GET()
 def do_POST(self):
  c=db();path=urlparse(self.path).path
  if path=='/api/auth/setup':
   if c.execute('SELECT COUNT(*) FROM users').fetchone()[0]:return self.send_json({'error':'管理员账号已创建'},409)
   try:data=json_body(self);username=str(data.get('username','')).strip();password=data.get('password','')
   except:return self.send_json({'error':'请求格式错误'},400)
   if len(username)<3 or len(username)>32 or not valid_password(password):return self.send_json({'error':'账号至少 3 位，密码至少 8 位'},400)
   salt,hashed=password_hash(password);c.execute('INSERT INTO users(username,password_hash,password_salt,created_at,role) VALUES(?,?,?,?,?)',(username,hashed,salt,datetime.now().isoformat(),'admin'));c.commit();return self.send_json({'ok':True})
  if path=='/api/auth/login':
   try:data=json_body(self);username=str(data.get('username','')).strip();password=data.get('password','')
   except:return self.send_json({'error':'请求格式错误'},400)
   key=client_key(self)+':'+username;now=time.time();attempts=[x for x in LOGIN_ATTEMPTS.get(key,[]) if now-x<900];LOGIN_ATTEMPTS[key]=attempts
   if len(attempts)>=5:return self.send_json({'error':'登录失败次数过多，请 15 分钟后再试'},429)
   user=c.execute('SELECT * FROM users WHERE username=?',(username,)).fetchone();ok=False
   if user:ok=secrets.compare_digest(password_hash(password,user['password_salt'])[1],user['password_hash'])
   if not ok:LOGIN_ATTEMPTS[key]=attempts+[now];return self.send_json({'error':'账号或密码错误'},401)
   token=secrets.token_urlsafe(32);expires=int(time.time()+SESSION_DAYS*86400);c.execute('INSERT INTO sessions(token_hash,user_id,expires_at,created_at) VALUES(?,?,?,?)',(hashlib.sha256(token.encode()).hexdigest(),user['id'],expires,datetime.now().isoformat()));c.commit();self.send_response(200);self.send_header('Content-Type','application/json; charset=utf-8');self.send_header('Set-Cookie',f'finance_session={token}; HttpOnly; SameSite=Lax; Max-Age={SESSION_DAYS*86400}; Path=/');self.send_header('Access-Control-Allow-Origin',self.headers.get('Origin','http://127.0.0.1:5173'));self.send_header('Access-Control-Allow-Credentials','true');self.end_headers();self.wfile.write(json.dumps({'ok':True,'user':{'username':user['username'],'role':user['role']}},ensure_ascii=False).encode());return
  if path=='/api/auth/logout':
   user=session_user(c,self);cookie=self.headers.get('Cookie','');token=next((x.split('=',1)[1] for x in cookie.split(';') if x.strip().startswith('finance_session=')), '')
   if token:c.execute('DELETE FROM sessions WHERE token_hash=?',(hashlib.sha256(token.encode()).hexdigest(),));c.commit()
   self.send_response(200);self.send_header('Set-Cookie','finance_session=; HttpOnly; SameSite=Lax; Max-Age=0; Path=/');self.send_header('Content-Type','application/json; charset=utf-8');self.end_headers();self.wfile.write(b'{"ok":true}');return
  if path=='/api/admin/delete-data':
   user=session_user(c,self)
   if not is_admin(user): return self.send_json({'error':'仅管理员可执行删除'},403)
   try:
    data=json_body(self); dates=[str(x).strip() for x in data.get('dates',[]) if str(x).strip()]; start=str(data.get('start_date','')).strip(); end=str(data.get('end_date','')).strip()
   except: return self.send_json({'error':'请求格式错误'},400)
   if dates:
    if any(len(x)!=10 for x in dates): return self.send_json({'error':'日期格式无效'},400)
    marks=','.join('?'*len(dates)); deleted=c.execute(f'DELETE FROM product_profit_daily WHERE stat_date IN ({marks})',dates).rowcount
   else:
    if not start or not end or len(start)!=10 or len(end)!=10 or start>end: return self.send_json({'error':'请选择有效的日期范围'},400)
    deleted=c.execute('DELETE FROM product_profit_daily WHERE stat_date>=? AND stat_date<=?',(start,end)).rowcount
   c.commit()
   return self.send_json({'ok':True,'deleted':deleted,'metrics':admin_metrics(c)})
  if path!='/api/import':self.send_error(404);return
  if not session_user(c,self):return self.send_json({'error':'请先登录'},401)
  try:
   length=int(self.headers.get('Content-Length','0')); body=self.rfile.read(length); envelope=(f"Content-Type: {self.headers.get('Content-Type','')}\r\nMIME-Version: 1.0\r\n\r\n").encode()+body; message=BytesParser(policy=policy.default).parsebytes(envelope); item=next((p for p in message.iter_parts() if p.get_param('name',header='Content-Disposition')=='file'),None)
   if item is None: raise ValueError('未找到上传文件')
   name=os.path.basename(item.get_filename() or 'import.xlsx');raw=item.get_payload(decode=True) or b'';h=hashlib.sha256(raw).hexdigest();values=[]
   if name.lower().endswith('.xlsx'):
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(raw),data_only=True,read_only=False)
    for ws in wb.worksheets:
     rows=list(ws.values); headers=[str(x or '').strip() for x in rows[0]]
     current_department=''
     current_manager=''
     current_promotion_leader=''
     current_date=''
     merged_rows={2:set(),3:set(),4:set()}
     for merged in ws.merged_cells.ranges:
      if merged.min_col==merged.max_col and merged.min_col in merged_rows:
       merged_rows[merged.min_col].update(range(merged.min_row,merged.max_row+1))
     for excel_row,r in enumerate(rows[1:],start=2):
      if len(r)<20:continue
      if r[0]: current_date=excel_date(r[0])
      product_name=str(r[5] or '').strip()
      if not current_date or not product_name or product_name in ('产品名称','商品名称') or any(x in product_name for x in ('合计','总计','小计')):continue
      if r[1]: current_department=str(r[1]).strip()
      if r[2]: current_manager=str(r[2]).strip()
      if r[3]: current_promotion_leader=str(r[3]).strip()
      department=current_department if (r[1] or excel_row in merged_rows[2]) else ''
      manager=current_manager if (r[2] or excel_row in merged_rows[3]) else ''
      promotion_leader=current_promotion_leader if (r[3] or excel_row in merged_rows[4]) else ''
      values.append((current_date,department,manager,promotion_leader,r[4] or '',product_name,*[num(x) for x in r[6:20]],ws.title,h))
    else:
     csv_text=raw.decode('utf-8-sig',errors='ignore').replace('\x00','')
     try:
      data=list(csv.reader(io.StringIO(csv_text, newline='')))
     except csv.Error as exc:
      raise ValueError('CSV 格式不规范：存在未加引号的换行或引号，请先用 Excel 另存为 XLSX 后再上传') from exc
     current_date=''
     for r in data[1:]:
      if len(r)<20:continue
      if r[0]:current_date=excel_date(r[0])
      product_name=str(r[5] or '').strip()
      if current_date and product_name and product_name not in ('产品名称','商品名称') and not any(x in product_name for x in ('合计','总计','小计')):values.append((current_date,r[1],r[2],r[3],r[4],product_name,*[num(x) for x in r[6:20]],'',h))
   c=db();now=datetime.now().isoformat();ins=upd=0
   key_cols=['stat_date','department','manager','promotion_leader','promotion_group','product_name']
   update_cols=COLS[6:]+['source_sheet','source_hash','updated_at']
   for row in values:
    row_values=dict(zip(COLS+['source_sheet','source_hash'],row)); key=tuple(row_values[k] for k in key_cols)
    existing=c.execute('SELECT id FROM product_profit_daily WHERE '+ ' AND '.join(f'{k}=?' for k in key_cols),key).fetchone()
    if existing:
     assignments=','.join(f'{k}=?' for k in update_cols); values_to_update=[row_values[k] for k in update_cols[:-1]]+[now,existing['id']]
     c.execute(f'UPDATE product_profit_daily SET {assignments} WHERE id=?',values_to_update);upd+=1
    else:
     full=row+(now,now);c.execute('INSERT INTO product_profit_daily ('+','.join(COLS)+',source_sheet,source_hash,created_at,updated_at) VALUES ('+','.join('?'*24)+')',full);ins+=1
   c.commit();return self.send_json({'inserted':ins,'updated':upd,'source_rows':len(values),'days':len(set(x[0] for x in values))})
  except Exception as e:return self.send_json({'error':str(e)},400)
if __name__=='__main__':
 c=db();sample_seed(c);port=int(os.environ.get('PORT','18780'));print(f'Finance API: http://127.0.0.1:{port}');ThreadingHTTPServer(('127.0.0.1',port),Handler).serve_forever()
